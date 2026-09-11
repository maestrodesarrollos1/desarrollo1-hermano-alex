from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .catalog import FieldDefinition, TemplateDefinition, discover_templates
from .workspace import TEMPLATES_ROOT


BRIEFING_DIRECTORY = TEMPLATES_ROOT.parent / "nupia_studio" / "briefings"
GLOBAL_BRIEFING_TEMPLATE_ID = "nupia-global"
PHOTO_SLOTS = (
    ("landing", "Portada", "Una foto vertical o panorámica de los dos. Es la imagen que abre la invitación."),
    ("historia1", "Historia 1", "Una imagen natural que cuente un momento compartido."),
    ("historia2", "Historia 2", "Un segundo momento con una composición distinta."),
    ("galeria1", "Galería 1", "Detalle editorial: manos, flores, tejidos o arquitectura."),
    ("galeria2", "Galería 2", "Retrato de la pareja o fotografía de ambiente."),
    ("galeria3", "Galería 3", "Un momento espontáneo, con espacio para recorte."),
    ("galeria4", "Galería 4", "Un detalle de la finca, mesa o paisaje."),
    ("galeria5", "Galería 5", "Una fotografía de contraste para dar ritmo a la galería."),
    ("galeria6", "Galería 6", "Cierre de la secuencia fotográfica."),
    ("despedida", "Despedida", "Imagen final cálida, preferiblemente vertical y con aire alrededor."),
)


class BriefingDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class BriefingImportResult:
    values: dict[str, Any]
    imported_fields: tuple[str, ...]
    template_id: str


def create_briefing_document(template: TemplateDefinition, destination: Path) -> Path:
    """Creates a client-ready Word briefing with stable keys for Studio import."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    _configure_styles(document)
    _add_header_footer(document, template.name)

    kicker = document.add_paragraph()
    kicker.style = "Nupia Kicker"
    is_global = template.template_id == GLOBAL_BRIEFING_TEMPLATE_ID
    kicker.add_run("NUPIA / BRIEFING UNIVERSAL DE BODA" if is_global else "NUPIA / BRIEFING DE BODA")

    title = document.add_paragraph()
    title.style = "Nupia Title"
    title.add_run("Toda vuestra boda, en un único documento" if is_global else f"Información para {template.name}")

    subtitle = document.add_paragraph()
    subtitle.style = "Nupia Intro"
    subtitle.add_run(
        "Rellena la columna de respuesta y guarda el archivo. Este briefing sirve para cualquier plantilla Nupia: "
        "elige primero el diseño en Studio, importa el Word y selecciona también la carpeta de fotos."
        if is_global
        else "Rellena la columna de respuesta y guarda el archivo. Después, en Nupia Studio, "
        "elige Importar briefing Word y selecciona también la carpeta de fotos."
    )

    marker = document.add_paragraph()
    marker.style = "Nupia Meta"
    marker.add_run(f"NUPIA_BRIEFING_TEMPLATE={template.template_id}")

    fields_by_group: dict[str, list[Any]] = {}
    for field in template.fields:
        if field.kind == "image":
            continue
        fields_by_group.setdefault(field.group, []).append(field)

    for group, fields in fields_by_group.items():
        document.add_heading(group, level=1)
        _add_instruction(document, group)
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_widths(table, (2.05, 3.55, 0.9))
        headers = ("Campo", "Respuesta", "Clave Studio")
        for cell, header in zip(table.rows[0].cells, headers):
            _set_cell_text(cell, header, bold=True, color="FFFFFF")
            _set_cell_fill(cell, "173E2F")
        for field in fields:
            row = table.add_row().cells
            _set_cell_text(row[0], field.label + (" *" if field.required else ""), bold=True)
            default = _format_default(field.default, field.kind)
            _set_cell_text(row[1], default)
            _set_cell_text(row[2], field.key, color="6D8477", size=8)
            for cell in row:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_table_widths(table, (2.05, 3.55, 0.9))
        document.add_paragraph().paragraph_format.space_after = Pt(5)

    document.add_heading("Fotos para pedir o seleccionar", level=1)
    photo_note = document.add_paragraph()
    photo_note.style = "Nupia Intro"
    photo_note.add_run(
        "Puedes aportar fotos definitivas o elegir fotos provisionales de stock con licencia. "
        "Colócalas todas en una carpeta y usa exactamente estos nombres. Studio reconocerá JPG, JPEG, PNG, WebP y AVIF."
    )
    photo_table = document.add_table(rows=1, cols=3)
    photo_table.style = "Table Grid"
    photo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_widths(photo_table, (1.45, 1.25, 3.8))
    for cell, header in zip(photo_table.rows[0].cells, ("Archivo", "Uso", "Foto que hay que pedir")):
        _set_cell_text(cell, header, bold=True, color="FFFFFF")
        _set_cell_fill(cell, "173E2F")
    supported_image_keys = {field.key.removeprefix("images.") for field in template.fields if field.kind == "image"}
    for slot, usage, direction in PHOTO_SLOTS:
        if slot not in supported_image_keys:
            continue
        row = photo_table.add_row().cells
        _set_cell_text(row[0], f"{slot}.jpg", bold=True)
        _set_cell_text(row[1], usage)
        _set_cell_text(row[2], direction)
    _set_table_widths(photo_table, (1.45, 1.25, 3.8))

    document.add_heading("Antes de enviarlo a Nupia", level=1)
    checklist = (
        "He revisado nombres, fecha, lugar y datos de contacto.",
        "He dejado las fotos en una sola carpeta con los nombres indicados.",
        "Las fotos de stock elegidas tienen licencia de uso y se pueden publicar.",
        "He guardado el Word sin cambiar la columna Clave Studio.",
    )
    for item in checklist:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def create_global_briefing_document(templates: list[TemplateDefinition], destination: Path) -> Path:
    if not templates:
        raise BriefingDocumentError("No hay plantillas disponibles para crear el briefing universal.")
    fields_by_key: dict[str, FieldDefinition] = {}
    for template in templates:
        for field in template.fields:
            fields_by_key.setdefault(field.key, field)
    ordered_fields = tuple(
        sorted(
            fields_by_key.values(),
            key=lambda field: (_global_group_order(field.group), field.label.casefold()),
        )
    )
    reference = templates[0]
    global_template = TemplateDefinition(
        template_id=GLOBAL_BRIEFING_TEMPLATE_ID,
        name="Briefing universal Nupia",
        description="Datos compartidos para cualquier plantilla Nupia.",
        descriptor_path=reference.descriptor_path,
        source_path=reference.source_path,
        preview_path=reference.preview_path,
        values_path=reference.values_path,
        fields=ordered_fields,
    )
    return create_briefing_document(global_template, destination)


def import_briefing_document(path: Path, template: TemplateDefinition, current_values: dict[str, Any]) -> BriefingImportResult:
    if path.suffix.lower() != ".docx":
        raise BriefingDocumentError("Selecciona un archivo Word .docx generado por Nupia.")
    try:
        document = Document(path)
    except Exception as exc:  # python-docx has no stable public exception hierarchy.
        raise BriefingDocumentError(f"No se pudo leer el documento Word: {exc}") from exc

    marker = next((paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.startswith("NUPIA_BRIEFING_TEMPLATE=")), "")
    template_id = marker.partition("=")[2].strip()
    if not template_id:
        raise BriefingDocumentError("Este Word no parece un briefing generado por Nupia Studio.")
    if template_id not in {GLOBAL_BRIEFING_TEMPLATE_ID, template.template_id}:
        raise BriefingDocumentError(
            f"El briefing corresponde a '{template_id}', pero la plantilla activa es '{template.template_id}'."
        )

    fields_by_key = {field.key: field for field in template.fields if field.kind != "image"}
    values = dict(current_values)
    imported: list[str] = []
    for table in document.tables:
        for row in table.rows[1:]:
            if len(row.cells) < 3:
                continue
            key = row.cells[2].text.strip()
            field = fields_by_key.get(key)
            if field is None:
                continue
            raw_value = row.cells[1].text.strip()
            if not raw_value:
                continue
            values[key] = _parse_value(raw_value, field.kind)
            imported.append(key)

    if not imported:
        raise BriefingDocumentError("No se encontraron respuestas importables. No modifiques la columna Clave Studio.")
    return BriefingImportResult(values, tuple(imported), template_id)


def generate_all_briefings(output_directory: Path = BRIEFING_DIRECTORY) -> list[Path]:
    documents: list[Path] = []
    templates = discover_templates(TEMPLATES_ROOT)
    documents.append(create_global_briefing_document(templates, output_directory / "briefing-universal-nupia.docx"))
    for template in templates:
        filename = f"briefing-{_slugify(template.name)}.docx"
        documents.append(create_briefing_document(template, output_directory / filename))
    return documents


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for style_name, size, color, before, after in (("Heading 1", 15, "173E2F", 16, 7), ("Heading 2", 12, "173E2F", 10, 5)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, size, color in (("Nupia Kicker", 8.5, "2E7D59"), ("Nupia Title", 26, "173E2F"), ("Nupia Intro", 10.5, "476456"), ("Nupia Meta", 7.5, "82968B")):
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(7)
    document.styles["Nupia Kicker"].font.bold = True
    document.styles["Nupia Kicker"].font.all_caps = True
    document.styles["Nupia Title"].font.name = "Georgia"
    document.styles["Nupia Title"].font.bold = False
    document.styles["Nupia Title"].paragraph_format.space_after = Pt(5)
    document.styles["Nupia Intro"].paragraph_format.line_spacing = 1.25


def _add_header_footer(document: Document, template_name: str) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(f"NUPIA  |  {template_name.upper()}")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("71877A")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Nupia Studio · briefing de contenidos")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("71877A")


def _add_instruction(document: Document, group: str) -> None:
    instructions = {
        "Datos básicos": "Escribe la información exactamente como debe verla una persona invitada.",
        "Textos": "Mantén un tono cercano y concreto. Puedes ampliar las celdas de respuesta cuanto necesites.",
        "Diseño": "Conserva los colores si no hay una dirección artística definida; puedes sustituirlos por códigos HEX.",
        "Contacto": "Incluye solo datos que puedan ser visibles para las personas invitadas.",
        "Secciones": "Responde Sí o No para mostrar u ocultar cada bloque.",
    }
    text = instructions.get(group, "Completa estos datos tal y como deben aparecer en la web.")
    paragraph = document.add_paragraph()
    paragraph.style = "Nupia Meta"
    paragraph.add_run(text)


def _format_default(value: Any, kind: str) -> str:
    if kind == "boolean":
        return "Sí" if bool(value) else "No"
    return str(value or "")


def _parse_value(value: str, kind: str) -> Any:
    if kind == "boolean":
        normalized = value.strip().lower()
        if normalized in {"si", "sí", "true", "1", "x", "mostrar"}:
            return True
        if normalized in {"no", "false", "0", "ocultar"}:
            return False
        raise BriefingDocumentError(f"'{value}' no es una respuesta válida. Usa Sí o No.")
    return value


def _set_table_widths(table: Any, widths: tuple[float, ...]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            _set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), "9360")
    table_width.set(qn("w:type"), "dxa")


def _set_cell_margins(cell: Any, *, top: int, bottom: int, start: int, end: int) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_fill(cell: Any, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_text(cell: Any, text: str, *, bold: bool = False, color: str = "173E2F", size: float = 9.5) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold


def _slugify(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "-", value.lower())
    return normalized.strip("-")


def _global_group_order(group: str) -> int:
    order = {"Datos básicos": 0, "Textos": 1, "Contacto": 2, "Diseño": 3, "Secciones": 4, "Fotos": 5, "Imágenes": 5}
    return order.get(group, 99)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Genera briefings Word de las plantillas Nupia")
    parser.add_argument("--output", type=Path, default=BRIEFING_DIRECTORY)
    args = parser.parse_args(argv)
    for path in generate_all_briefings(args.output):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
